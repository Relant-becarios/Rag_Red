"""Excel tables (.xlsx / .xls) as product sources + datasheets of a part number that is in
several rows (the 409 Conflict that broke the PDF download link in the chat).

    python -m pytest tests/test_excel_datasheet.py -q
"""
import shutil
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import test_products as tp  # noqa: E402  (test environment + FakeEmb + sample data)


def _own_folders():
    """Own temp folder: test_products deletes its folder when it finishes."""
    import tempfile
    root = Path(tempfile.mkdtemp(prefix="relant_excel_test_"))
    tp.DATA, tp.INDEX = root / "sample_data", root / "sample_index"
    return root


def make_excel():
    import openpyxl
    d = tp.DATA / "productos"
    # 1) two sheets, a title row above the header, a merged cell, numbers as Excel shows them
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Neumaticos"
    ws.append(["Catálogo de actuadores 2026"])
    ws.append(["Número de parte", "Nombre del producto", "Carrera", "Fuerza", "Eficiencia",
               "Serie"])
    ws.append([4192401, "Actuador neumático A1", 12.0, 0.1 + 0.2, 0.85, "Serie X"])
    ws.append(["AX-4193000", "Válvula AX-4193000", 25, 150.5, 0.9, None])
    ws["E3"].number_format = "0%"
    ws["E4"].number_format = "0%"
    ws.merge_cells("F3:F4")                       # "Serie X" for both rows
    ws2 = wb.create_sheet("Electricos")
    ws2.append(["Número de parte", "Nombre del producto", "Voltaje", "Carrera"])
    ws2.append(["EL-50", "Actuador eléctrico EL-50", "24 V", 50])
    hidden = wb.create_sheet("interno")
    hidden.append(["Número de parte", "Costo secreto"])
    hidden.append(["EL-50", 999])
    hidden.sheet_state = "hidden"
    wb.save(d / "actuadores.xlsx")

    # 2) the same part number in another table with DIFFERENT data -> variant column
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Número de parte", "Nombre del producto", "Carrera", "Material"])
    ws.append(["AX-4193000", "Válvula AX-4193000", 30, "Acero inoxidable"])
    wb.save(d / "valvulas.xlsx")

    # 3) bombas.csv exported again as .xlsx: identical rows must NOT become a conflict
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Número de parte", "Nombre del producto", "Presión máx.", "Caudal",
               "Conexión de salida", "Peso", "Notas internas"])
    ws.append(["BP-200", "Bomba de pistón BP-200", "20-220 bar", "1,2 l/min", 'G 1/8"',
               "3,5 kg", "no mostrar"])
    wb.save(d / "bombas_copia.xlsx")

    # 4) old Excel 97-2003 .xls
    import xlwt
    book = xlwt.Workbook()
    sh = book.add_sheet("Hoja1")
    for j, h in enumerate(["No. de parte", "Producto", "Presión"]):
        sh.write(0, j, h)
    for j, v in enumerate(["OLD-1", "Válvula antigua OLD-1", 6.0]):
        sh.write(1, j, v)
    book.save(str(d / "antiguos.xls"))
    (d / "~$actuadores.xlsx").write_bytes(b"lock file of an open workbook")


def main():
    tp.make_data()
    make_excel()
    from app.products import build_datasheet_html, build_datasheet_pdf, get_product_index
    from app.retriever import HybridEvaluationRetriever
    from app.textutils import read_csv_table

    h, rows, pre = read_csv_table(tp.DATA / "productos/actuadores.xlsx")
    assert h[0] == "Hoja" and "Voltaje" in h and "Costo secreto" not in h, h
    assert pre == ["Neumaticos: Catálogo de actuadores 2026"], pre
    r1 = dict(zip(h, rows[0]))
    assert r1["Número de parte"] == "4192401" and r1["Carrera"] == "12", r1
    assert r1["Fuerza"] == "0.3" and r1["Eficiencia"] == "85%", r1
    assert dict(zip(h, rows[1]))["Serie"] == "Serie X"           # merged cell filled
    assert dict(zip(h, rows[2]))["Hoja"] == "Electricos"

    r = HybridEvaluationRetriever(data_dir=str(tp.DATA), chroma_dir=str(tp.INDEX / "chroma"),
                                  manifest_file=str(tp.INDEX / "manifest.json"),
                                  embeddings=tp.FakeEmb())
    r.update_index()
    cat = r.snapshot.catalog
    for rel in ("productos/actuadores.xlsx", "productos/valvulas.xlsx", "productos/antiguos.xls"):
        assert cat[rel]["type"] == "table", rel
    assert not any("~$" in k for k in cat)
    assert "4192401" in r.table_lookup("4192401")

    idx = get_product_index(r)
    assert idx.find("EL-50")[0].name == "Actuador eléctrico EL-50"
    assert idx.find("OLD-1")[0].name == "Válvula antigua OLD-1"
    assert idx.find("4192401")

    # identical rows (bombas.csv + bombas_copia.xlsx): one datasheet, no variants
    p, variants, cands = idx.resolve("BP-200")
    assert p is not None and variants == [] and cands == []
    assert p.image is not None

    # same part, different data: one datasheet with two columns
    p, variants, cands = idx.resolve("AX-4193000")
    assert p is not None and len(variants) == 1, (p, variants, cands)
    page = build_datasheet_html(p, variants=variants)
    assert "Acero inoxidable" in page and "150.5" in page
    pdf = build_datasheet_pdf(p, variants)
    assert pdf[:5] == b"%PDF-"
    import pymupdf as fitz
    with fitz.open(stream=pdf, filetype="pdf") as doc:
        text = "".join(pg.get_text() for pg in doc)
    assert "Acero inoxidable" in text and "150.5" in text, text

    # different products sharing a word: ambiguous -> candidates, no product
    p, variants, cands = idx.resolve("Actuador")
    assert p is None and len(cands) >= 2

    # the chat tool collapses the duplicates and offers ONE working download link
    from app.tools import RequestContext, make_tools
    ctx = RequestContext(retriever=r)
    out = make_tools(r, [], ctx)["product_datasheet"](part="AX-4193000")
    assert "variant" in out and len(ctx.downloads) == 1, (out, ctx.downloads)
    print("OK")


def test_excel_and_duplicate_parts() -> None:
    saved = tp.DATA, tp.INDEX
    root = _own_folders()
    try:
        main()
    finally:
        tp.DATA, tp.INDEX = saved
        shutil.rmtree(root, ignore_errors=True)


if __name__ == "__main__":
    test_excel_and_duplicate_parts()


def test_datasheet_endpoints() -> None:
    """/datasheet-pdf/{part}: 200 + a real PDF for a repeated part number (was 409)."""
    import test_users as tu
    from fastapi.testclient import TestClient
    from app.config import settings
    from app.db import reset_engine
    from app.jobs import IndexJobManager
    from app.main import create_app
    from app.retriever import HybridEvaluationRetriever
    from app.router import QueryRouter

    saved = tp.DATA, tp.INDEX
    root = _own_folders()
    try:
        tp.make_data()
        make_excel()
        if not tu.os.getenv("TEST_DATABASE_URL"):
            settings.database_url = f"sqlite:///{(root / 'app.db').as_posix()}"
        reset_engine()
        tu.base.reset_database()
        r = HybridEvaluationRetriever(data_dir=str(tp.DATA), chroma_dir=str(tp.INDEX / "chroma"),
                                      manifest_file=str(tp.INDEX / "manifest.json"),
                                      embeddings=tp.FakeEmb())
        r.update_index()
        comps = {"retriever": r, "router": QueryRouter(r, tu.base.FakeLLM()),
                 "jobs": IndexJobManager(r, run_preflight=False)}
        tu.base.make_user("root", "admin")
        app = create_app(comps, auto_sync=False, web_warmup=False)
        with TestClient(app) as c:
            tu.base.login(c, "root")
            res = c.get("/datasheet-pdf/AX-4193000")
            assert res.status_code == 200, res.text
            assert res.headers["content-type"] == "application/pdf"
            assert res.content[:5] == b"%PDF-"
            assert "attachment" in res.headers["content-disposition"]
            assert c.get("/datasheet-pdf/BP-200").status_code == 200
            assert c.get("/datasheet/AX-4193000").status_code == 200
            # a base that does not exist / is not readable does not break the link
            assert c.get("/datasheet-pdf/AX-4193000?kb=general,u-otra-persona").status_code == 200
            # errors: a PDF with the reason (never a JSON file saved as <part>.json)
            from urllib.parse import unquote
            amb = c.get("/datasheet-pdf/Actuador")
            assert amb.status_code == 409 and amb.content[:5] == b"%PDF-"
            assert "número de parte" in unquote(amb.headers["x-error-detail"])
            nf = c.get("/datasheet-pdf/NO-EXISTE-1")
            assert nf.status_code == 404 and nf.content[:5] == b"%PDF-"
            assert c.get("/datasheet/NO-EXISTE-1").headers["content-type"].startswith("text/html")
            assert c.get("/healthz").json()["build"].startswith("2026-10-19b")
    finally:
        tp.DATA, tp.INDEX = saved
        shutil.rmtree(root, ignore_errors=True)
